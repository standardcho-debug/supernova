"""layers.json → Canva 가져오기용 PPTX 2종 (캐러셀 5p, 광고 1p). 1080×1350px, 텍스트는 전부 텍스트박스.
실행: python3 extract.py && python3 build_pptx.py"""
import json, os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PX = 9525  # 1px = 9525 EMU (96dpi)
FONT = 'Pretendard'
D = json.load(open(os.path.join(HERE, 'layers.json')))


def px(v): return Emu(int(round(v * PX)))
def rgb(h): return RGBColor.from_string(h.lstrip('#'))


def ad_composite():
    """광고 하단 지도: 둥근 모서리 클리핑 + 크림색 페이드를 미리 합성한 이미지 1장."""
    a = D['ad']['adshot']; src = Image.open(os.path.join(ROOT, 'src/map_crop.png')).convert('RGBA')
    k = src.width / a['iw']; W, H = round(a['w'] * k), round(a['h'] * k)
    box = Image.new('RGBA', (W, H), (250, 248, 243, 255))
    img = src.resize((round(a['iw'] * k), round(a['ih'] * k)))
    box.paste(img, (round(a['ix'] * k), round(a['iy'] * k)), img)
    fade = Image.new('RGBA', (W, H), (0, 0, 0, 0)); d = ImageDraw.Draw(fade)
    y0 = int(H * .55)
    for y in range(y0, H):
        d.line([(0, y), (W, y)], fill=(250, 248, 243, round(255 * (y - y0) / (H - y0))))
    box.alpha_composite(fade)
    mask = Image.new('L', (W, H), 0); ImageDraw.Draw(mask).rounded_rectangle([0, 0, W - 1, H - 1], radius=round(24 * k), fill=255)
    box.putalpha(mask)
    out = os.path.join(HERE, 'ad_map_composite.png'); box.save(out); return out


def add_shape(sl, L):
    if L.get('ellipse'):
        sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, px(L['x']), px(L['y']), px(L['w']), px(L['h']))
    elif L.get('radius', 0) > 0.5:
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(L['x']), px(L['y']), px(L['w']), px(L['h']))
        sh.adjustments[0] = min(0.5, L['radius'] / min(L['w'], L['h']))
    else:
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(L['x']), px(L['y']), px(L['w']), px(L['h']))
    sh.shadow.inherit = False
    if L.get('fill'): sh.fill.solid(); sh.fill.fore_color.rgb = rgb(L['fill'])
    else: sh.fill.background()
    if L.get('line'):
        sh.line.color.rgb = rgb(L['line']); sh.line.width = Pt(L['lineW'] * .75)
        if L.get('dash'): sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: sh.line.fill.background()
    sh.name = L['name']
    sh.text_frame.text = ''
    return sh


def container(L, shapes):
    """작은 채움 도형 안에 완전히 들어간 텍스트 → 도형 영역에 가운데 정렬(편집 시에도 중앙 유지)."""
    for s in shapes:
        if not s.get('fill') or s.get('ellipse') and s['w'] > 200: continue
        if s['w'] > 400 or s['h'] > 130: continue
        if s['x'] - 1 <= L['x'] and s['y'] - 1 <= L['y'] and L['x'] + L['w'] <= s['x'] + s['w'] + 1 and L['y'] + L['h'] <= s['y'] + s['h'] + 1:
            return s
    return None


def add_text(sl, L, shapes):
    c = container(L, shapes)
    if c: x, y, w, h = c['x'], c['y'], c['w'], c['h']
    else: x, y, w, h = L['x'], L['y'], L['w'] * 1.06 + 4, L['h']
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h)); tb.name = 'T ' + L['name']
    tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if c else MSO_ANCHOR.TOP
    for i, line in enumerate(L['lines']):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
        if L.get('lh') and len(L['lines']) > 1: p.line_spacing = Pt(L['lh'] * .75)
        for r in line:
            run = p.add_run(); run.text = r['t']; f = run.font
            f.size = Pt(r['size'] * .75); f.bold = r['weight'] >= 600; f.color.rgb = rgb(r['color']); f.name = FONT
            rPr = run._r.get_or_add_rPr()
            for tag in ('a:ea', 'a:cs'):
                el = rPr.find(qn(tag))
                if el is None: el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set('typeface', FONT)
            if r['ls']: rPr.set('spc', str(int(round(r['ls'] * .75 * 100))))
            if r['strike']: rPr.set('strike', 'sngStrike')
    return tb


def build(ids, fname, title):
    prs = Presentation(); prs.slide_width, prs.slide_height = px(1080), px(1350)
    prs.core_properties.title = title
    comp = ad_composite() if 'ad' in ids else None
    for sid in ids:
        S = D[sid]; sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(S['bg'])
        shapes = [L for L in S['layers'] if L['kind'] == 'shape']
        for L in S['layers']:  # DOM 순서 = 쌓임 순서
            if sid == 'ad' and L['kind'] == 'shape' and L['name'] == 'adshot':
                a = S['adshot']; sl.shapes.add_picture(comp, px(a['x']), px(a['y']), px(a['w']), px(a['h'])).name = 'IMG 결정지도(합성)'
                continue
            if sid == 'ad' and L['kind'] == 'image' and 'map_crop' in L['src']: continue
            if L['kind'] == 'shape': add_shape(sl, L)
            elif L['kind'] == 'image':
                sl.shapes.add_picture(os.path.join(ROOT, L['src']), px(L['x']), px(L['y']), px(L['w']), px(L['h'])).name = 'IMG ' + os.path.basename(L['src'])
            else: add_text(sl, L, shapes)
    out = os.path.join(HERE, fname); prs.save(out); print('saved', out)


build(['s1', 's2', 's3', 's4', 's5'], 'TIPS_캐러셀_5p.pptx', 'TIPS 캐러셀 v1 (편집본)')
build(['ad'], 'TIPS_광고_피드_4x5.pptx', 'TIPS 피드 광고 v1 (편집본)')
